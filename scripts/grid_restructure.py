#!/usr/bin/env python3
"""
Restructure the comparison-grid tabs so the appraiser is never fighting the
auto comp picker.

For every sale column on a grid we now keep THREE layers:

  * MAIN grid  (unchanged location) - the single source of truth used by all
    downstream adjustments, the value conclusion and the report. Each comp
    cell is `=IF(<source toggle>="Manual", <manual cell>, <auto cell>)`, so it
    always shows a clean value and never has to be edited to override a comp.
  * AUTO-SELECTED block - what the comp picker fed in (the original formulas,
    e.g. links to Export030823). Read-only reference.
  * MANUAL OVERRIDE block - blank cells the appraiser types into. Typing here
    can never break a formula or an automation.

A per-sale Source toggle (Auto / Manual) sits above each sale column. Leaving
it on "Auto" loads the picker's comp; flipping it to "Manual" uses the typed
override. That toggle is the macro-free equivalent of a per-sale load button.
"""
import shutil
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from copy import copy

SRC = "Analysis_Workbook_SandboxCleanFinal.xlsx"

YELLOW = PatternFill("solid", fgColor="FFF2CC")     # manual input
GREY = PatternFill("solid", fgColor="F2F2F2")       # auto / reference
HDR = PatternFill("solid", fgColor="4472C4")
TOGGLE = PatternFill("solid", fgColor="E2EFDA")     # source toggle
HDR_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
NOTE = Font(italic=True, size=9, color="595959")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def transform(ws, cfg):
    label_col = cfg["label_col"]
    sale_cols = cfg["sale_cols"]
    raw_rows = cfg["raw_rows"]
    toggle_row = cfg["toggle_row"]
    header_row = cfg["header_row"]
    feed_hdr = cfg["feed_base"]
    man_hdr = cfg["manual_base"]
    feed_mode = cfg.get("feed_mode", "move")     # "move" or "export"
    export_map = cfg.get("export_map", {})       # grid_row -> Export030823 column
    export_row0 = cfg.get("export_row0", 3)      # Export row of Sale 1
    feed0 = feed_hdr + 2      # first data row of auto block
    man0 = man_hdr + 2        # first data row of manual block

    # ---- section banners --------------------------------------------------
    feed_title = cfg.get(
        "feed_title",
        "AUTO-SELECTED COMPS  (what the picker fed in - reference only)")
    ws.cell(feed_hdr, column=_ci(label_col), value=feed_title)
    ws.cell(feed_hdr, column=_ci(label_col)).font = HDR_FONT
    ws.cell(feed_hdr, column=_ci(label_col)).fill = HDR
    ws.cell(man_hdr, column=_ci(label_col),
            value="MANUAL OVERRIDES  (type here; used when a sale's Source = Manual)")
    ws.cell(man_hdr, column=_ci(label_col)).font = HDR_FONT
    ws.cell(man_hdr, column=_ci(label_col)).fill = HDR

    # echo the sale headers on both blocks (use "Sale n" to avoid merged cells)
    for blk in (feed_hdr + 1, man_hdr + 1):
        for n, sc in enumerate(sale_cols, start=1):
            dst = ws[f"{sc}{blk}"]
            dst.value = f"Sale {n}"
            dst.font = BOLD
            dst.alignment = Alignment(horizontal="center")

    # ---- per raw attribute row -------------------------------------------
    for i, r in enumerate(raw_rows):
        feed_r = feed0 + i
        man_r = man0 + i
        # carry the row label into both blocks
        lbl = ws[f"{label_col}{r}"].value
        ws[f"{label_col}{feed_r}"] = lbl
        ws[f"{label_col}{man_r}"] = lbl
        for c in (f"{label_col}{feed_r}", f"{label_col}{man_r}"):
            ws[c].font = Font(size=9)

        for n, sc in enumerate(sale_cols, start=1):
            main = ws[f"{sc}{r}"]
            numfmt = main.number_format
            if feed_mode == "export" and r in export_map:
                original = f"=Export030823!{export_map[r]}{export_row0 + n - 1}"
            else:
                original = main.value                   # picker formula / value

            feed = ws[f"{sc}{feed_r}"]
            man = ws[f"{sc}{man_r}"]
            feed.value = original
            feed.number_format = numfmt
            feed.fill = GREY
            feed.border = BOX
            man.value = None
            man.number_format = numfmt
            man.fill = YELLOW
            man.border = BOX

            # MAIN becomes the composed source-of-truth cell
            main.value = (f'=IF({sc}${toggle_row}="Manual",'
                          f'{sc}{man_r},{sc}{feed_r})')
            main.number_format = numfmt

    # ---- per-sale source toggles above each sale column -------------------
    first_sale = min(_ci(sc) for sc in sale_cols)
    # free up the toggle row: unmerge anything crossing it, keep the title
    # text in its left-most cell, then re-merge just the label area for it.
    title_val = None
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row <= toggle_row <= rng.max_row:
            tl = ws.cell(rng.min_row, rng.min_col)
            if title_val is None:
                title_val = tl.value
            ws.unmerge_cells(str(rng))
    if title_val is not None and first_sale - 1 > _ci(label_col):
        ws.merge_cells(start_row=toggle_row, start_column=_ci(label_col),
                       end_row=toggle_row, end_column=first_sale - 1)
        ws.cell(toggle_row, _ci(label_col)).value = title_val

    dv = DataValidation(type="list", formula1='"Auto,Manual"', allow_blank=False,
                        showInputMessage=True)
    dv.promptTitle = "Comp source"
    dv.prompt = ("Auto = use the comp the picker selected. Manual = use the value "
                 "you type in the MANUAL OVERRIDES block below. Nothing breaks "
                 "either way.")
    for sc in sale_cols:
        t = ws[f"{sc}{toggle_row}"]
        t.value = "Auto"
        t.fill = TOGGLE
        t.border = BOX
        t.alignment = Alignment(horizontal="center")
        t.font = BOLD
        dv.add(f"{sc}{toggle_row}")
    ws.add_data_validation(dv)

    # a short how-to note under the manual block
    note_r = man0 + len(raw_rows) + 1
    ws.cell(note_r, column=_ci(label_col),
            value=("Tip: keep Source on Auto to use the picker's comp. Switch a "
                   "sale to Manual and type its column here to override - the main "
                   "grid, adjustments and conclusion update with no formula edits."))
    ws.cell(note_r, column=_ci(label_col)).font = NOTE


def _ci(letter):
    from openpyxl.utils import column_index_from_string
    return column_index_from_string(letter)


CONFIGS = {
    "Basic Sales Grid": {
        "label_col": "C",
        "sale_cols": ["E", "F", "G", "H", "I"],
        "header_row": 3,
        "toggle_row": 2,
        "raw_rows": [4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14,
                     19, 20, 21, 22, 23, 28, 29, 30],
        "feed_base": 73,
        "manual_base": 96,
        "feed_mode": "move",
    },
    # Detailed grid is an empty template: wire the AUTO block to the same
    # Export030823 picker feed for the descriptive comp rows.
    "Detailed Sales Grid": {
        "label_col": "C",
        "sale_cols": ["E", "G", "I", "K", "M"],
        "header_row": 3,
        "toggle_row": 2,
        "raw_rows": [4, 5, 6, 7, 8, 9, 10, 11, 13, 14,
                     19, 21, 23, 25, 27, 32, 39, 40, 41, 47, 53],
        "feed_base": 72,
        "manual_base": 96,
        "feed_mode": "export",
        "export_row0": 3,
        "export_map": {4: "B", 5: "C", 6: "D", 7: "E", 8: "F", 9: "G",
                       10: "I", 11: "J", 13: "K", 14: "L", 19: "M", 21: "N",
                       23: "O", 25: "P", 27: "Q", 32: "R", 39: "S", 40: "T",
                       41: "U", 47: "W", 53: "X"},
    },
    # Land comps are hand-entered (no external picker): the AUTO block holds the
    # currently selected land sales; MANUAL lets the appraiser override safely.
    "Land Sales": {
        "label_col": "B",
        "sale_cols": ["E", "G", "I", "K", "L"],
        "header_row": 8,
        "toggle_row": 7,
        "raw_rows": [9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21, 22],
        "feed_base": 70,
        "manual_base": 90,
        "feed_mode": "move",
        "feed_title": ("SELECTED LAND COMPS  (currently loaded comps - reference; "
                       "no external picker for land)"),
    },
}


def build():
    shutil.copyfile(SRC, SRC + ".bak2")
    wb = openpyxl.load_workbook(SRC)
    for sheet, cfg in CONFIGS.items():
        transform(wb[sheet], cfg)
    try:
        wb.calculation.fullCalcOnLoad = True
    except Exception:
        pass
    wb.save(SRC)
    print("saved", SRC)


if __name__ == "__main__":
    build()
