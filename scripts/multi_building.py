#!/usr/bin/env python3
"""
Multi-building set-up for the CRE valuation workbook.

Adds a proper multi-building capture flow to 'Property Info':
  * "Number of Buildings" is entered after the property address.
  * Each building gets its own input column to the right (Building 1 = col C,
    Building 2 = col D, ... up to Building 12 = col N). Columns "appear"
    (get highlighted) as the building count increases.
  * The Subject / Total column (B) is auto-derived from the buildings
    (total GBA/NRA/footprint, GBA-weighted office %, and the primary use /
    year / condition of the largest building).

Feeds the Cost Approach: each building carries its own occupancy code, class,
quality, area, height, area/perimeter factors, sprinkler tier and effective
age. The Cost Approach now sums a per-building Replacement-Cost-New and blends
physical depreciation across buildings.

Feeds the Sales Approach: total square footage and primary use flow from the
Subject / Total column into the Basic Sales Grid subject.
"""
import re
import shutil
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

SRC = "Analysis_Workbook_SandboxCleanFinal.xlsx"

# ---- styling conventions pulled from the existing workbook -------------------
YELLOW = PatternFill("solid", fgColor="FFFF00")          # Property Info inputs
SOFTYEL = PatternFill("solid", fgColor="FFF2CC")         # Cost Approach inputs
HDR_BLUE = PatternFill("solid", fgColor="4472C4")        # section header
SUBHDR = PatternFill("solid", fgColor="D9E1F2")          # sub header / auto rows
AUTO_GREY = PatternFill("solid", fgColor="F2F2F2")       # auto (calculated) cells
WHITE_BOLD = Font(bold=True, color="FFFFFF")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
BOLD = Font(bold=True)
NOTE_FONT = Font(italic=True, size=9, color="595959")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

N_BLDG = 12                                 # buildings supported (cols C..N)
BLDG_COLS = [get_column_letter(2 + k) for k in range(1, N_BLDG + 1)]   # C..N

# rows on Property Info that are captured per-building
PHYS_ROWS = {
    2: "account", 5: "gba", 6: "footprint", 7: "nra",
    9: "yearbuilt", 10: "condition", 15: "officepct", 16: "use",
}
# new per-building COST rows appended below the subject block
COST_ROWS = {
    35: "Occupancy Type (MVS)",
    36: "Building Class",
    37: "Quality",
    38: "Avg Eave / Wall Height (ft)",
    39: "Area/Perimeter Factor 1",
    40: "Area/Perimeter Factor 2",
    41: "Sprinkler Tier",
    42: "Effective Age (yrs)",
}
ALL_INPUT_ROWS = sorted(list(PHYS_ROWS) + list(COST_ROWS))


def sub_refs(formula, mapping):
    """Replace whole $C$<n> tokens (not followed by another digit)."""
    out = formula
    # longmost-specific first, and use negative-lookahead on trailing digit
    for token, repl in sorted(mapping.items(), key=lambda kv: -len(kv[0])):
        num = token.replace("$C$", "")
        out = re.sub(r"\$C\$" + num + r"(?!\d)", repl, out)
    return out


def build():
    shutil.copyfile(SRC, SRC + ".bak")
    wb = openpyxl.load_workbook(SRC)

    property_info(wb)
    cost_approach(wb)
    sales_grid(wb)

    # formulas are written without cached values -> force a recalc on open
    try:
        wb.calculation.fullCalcOnLoad = True
    except Exception:
        pass

    wb.save(SRC)
    print("saved", SRC)


# =============================================================================
# 1. PROPERTY INFO
# =============================================================================
def property_info(wb):
    ws = wb["Property Info"]

    # ---- row 1: header labels ------------------------------------------------
    ws["A1"] = "Field"
    ws["A1"].fill = SUBHDR
    ws["A1"].font = BOLD
    ws["B1"] = "Subject / Total"
    ws["B1"].fill = SUBHDR
    ws["B1"].font = BOLD
    ws["B1"].alignment = Alignment(horizontal="center")
    for k, col in enumerate(BLDG_COLS, start=1):
        c = ws[f"{col}1"]
        c.value = f'=IF($B$4>={k},"Building {k}","")'
        c.fill = SUBHDR
        c.font = BOLD
        c.alignment = Alignment(horizontal="center")

    # ---- Number of Buildings: mirror at B4, real input after the address -----
    ws["A4"] = "Number of Buildings (auto — set in Multi-Building Setup below)"
    ws["A4"].font = NOTE_FONT
    ws["B4"] = "=B33"                       # keep B4 stable for downstream refs
    ws["B4"].fill = PatternFill(fill_type=None)

    # ---- Subject / Total column becomes auto-derived from the buildings ------
    def rng(row):
        return f"C{row}:{BLDG_COLS[-1]}{row}"

    ws["B5"] = f"=SUM({rng(5)})"                                   # total GBA
    ws["B6"] = f"=SUM({rng(6)})"                                   # total footprint
    ws["B7"] = f"=SUM({rng(7)})"                                   # total NRA
    ws["B9"] = f'=IFERROR(INDEX({rng(9)},MATCH(MAX({rng(5)}),{rng(5)},0)),"")'
    ws["B10"] = f'=IFERROR(INDEX({rng(10)},MATCH(MAX({rng(5)}),{rng(5)},0)),"")'
    ws["B11"] = '=IF(B9="","",B9&"/"&B10)'
    ws["B15"] = f'=IFERROR(SUMPRODUCT({rng(5)},{rng(15)})/SUM({rng(5)}),"")'
    ws["B16"] = f'=IFERROR(INDEX({rng(16)},MATCH(MAX({rng(5)}),{rng(5)},0)),"")'  # primary use
    # mark the now-auto subject cells so the user knows not to type over them
    for r in (5, 6, 7, 9, 10, 11, 15, 16):
        ws[f"B{r}"].fill = AUTO_GREY

    # ---- per-building physical inputs (existing rows, building columns) ------
    for row in PHYS_ROWS:
        numfmt = None
        if row in (5, 6, 7):
            numfmt = "#,##0"
        elif row == 15:
            numfmt = "0.0%"
        for col in BLDG_COLS:
            c = ws[f"{col}{row}"]
            c.border = BOX
            if numfmt:
                c.number_format = numfmt

    # ---- Multi-Building Setup panel (after the address block) ---------------
    ws["A32"] = "MULTI-BUILDING SETUP"
    ws["A32"].fill = HDR_BLUE
    ws["A32"].font = HDR_FONT
    for col in ["B"] + BLDG_COLS:
        ws[f"{col}32"].fill = HDR_BLUE

    ws["A33"] = "Number of Buildings"
    ws["A33"].font = BOLD
    ws["B33"] = 1
    ws["B33"].fill = YELLOW
    ws["B33"].border = BOX
    ws["B33"].alignment = Alignment(horizontal="center")
    ws["C33"] = ("Enter the count here, then fill one column per building to the "
                 "right (Building 1 = column C). Columns highlight as you raise the count.")
    ws["C33"].font = NOTE_FONT

    # cost-input row labels + building cells
    for row, label in COST_ROWS.items():
        ws[f"A{row}"] = label
        for col in BLDG_COLS:
            c = ws[f"{col}{row}"]
            c.border = BOX
            if row in (38, 39, 40, 42):
                c.number_format = "General"

    # Building 1 keeps the Cost Approach template defaults (so the sheet still
    # produces a value out of the box; the count defaults to 1).
    b1_defaults = {35: "Storage Warehouses", 36: "C", 37: "Average", 38: 18,
                   39: 1, 40: 1, 41: "Average", 42: 0}
    for row, val in b1_defaults.items():
        ws[f"C{row}"] = val

    # ---- data validations on the building columns ---------------------------
    col_span = f"C{{r}}:{BLDG_COLS[-1]}{{r}}"

    def add_dv(row, formula1, allow_blank=True, kind="list"):
        dv = DataValidation(type=kind, formula1=formula1, allow_blank=allow_blank)
        dv.add(col_span.format(r=row))
        ws.add_data_validation(dv)

    add_dv(35, "=OCC_NAME")
    add_dv(36, '"A,B,C,CMILL,D,DPOLE,S,A-B"')
    add_dv(37, '"Low Cost,Average,Good,Excellent,Cheap"')
    add_dv(41, '"Low Cost,Average,Good,Excellent,Cheap"')
    add_dv(10, '"Excellent,Good,Average,Fair,Poor"')
    dv_cnt = DataValidation(type="whole", operator="between",
                            formula1="1", formula2=str(N_BLDG), allow_blank=False)
    dv_cnt.add("B33")
    ws.add_data_validation(dv_cnt)

    # ---- conditional formatting: reveal / highlight active building columns --
    # (a) highlight active building input cells yellow;
    # (b) hide inactive columns' content with white font (matches the workbook's
    #     existing LEN(header)=0 rule, extended to col N and the new cost rows).
    white_hide = Font(color="FFFFFF")
    for k, col in enumerate(BLDG_COLS, start=1):
        cells = " ".join(f"{col}{r}" for r in ALL_INPUT_ROWS)
        ws.conditional_formatting.add(
            cells,
            FormulaRule(formula=[f"$B$4>={k}"], fill=YELLOW, stopIfTrue=False),
        )
        hide_cells = " ".join(f"{col}{r}" for r in ([1] + ALL_INPUT_ROWS))
        ws.conditional_formatting.add(
            hide_cells,
            FormulaRule(formula=[f"LEN(${col}$1)=0"], font=white_hide, stopIfTrue=False),
        )

    # widen building columns a touch for readability
    for col in BLDG_COLS:
        ws.column_dimensions[col].width = 15


# =============================================================================
# 2. COST APPROACH
# =============================================================================
def cost_approach(wb):
    ws = wb["Cost Approach"]

    # ---- Building 1 inputs now read from Property Info Building 1 (col C) -----
    ws["C4"] = "='Property Info'!C35"      # occupancy
    ws["C6"] = "='Property Info'!C36"      # class
    ws["C7"] = "='Property Info'!C37"      # quality
    ws["C11"] = "='Property Info'!C38"     # height
    ws["C12"] = "='Property Info'!C39"     # area/perim 1
    ws["C13"] = "='Property Info'!C40"     # area/perim 2
    ws["C14"] = "='Property Info'!C41"     # sprinkler tier
    ws["C24"] = "='Property Info'!C5"      # Building-1 area (GBA)
    ws["C58"] = "='Property Info'!C42"     # Building-1 effective age
    ws["I24"] = ('=IFERROR("Verified MVS base cost -- "&INDEX(BC_CITE,'
                 'MATCH($C$4&"~"&$C$6&"~"&$C$7,BC_KEY,0)),"")'
                 '&" | Area auto-linked to Property Info Building 1 (GBA)."')
    # Steps 1-3 now come from Property Info; tell the user
    ws["I4"] = ("Building-1 identity now flows from Property Info (Building 1 column). "
                "Add more buildings there; buildings 2-12 total in below.")

    # per-square-foot denominators use TOTAL GBA, not Building-1 area
    ws["G52"] = "=caRCN/'Property Info'!B5"
    ws["G65"] = "=caRCNDeprec/'Property Info'!B5"
    ws["G71"] = "=G70/'Property Info'!B5"

    # ---- template formulas from the existing single-building calc ------------
    T = {c[0].coordinate: c[0].value
         for c in [[ws["D24"]], [ws["D25"]], [ws["D33"]], [ws["D34"]],
                   [ws["D36"]], [ws["C59"]], [ws["C60"]]]}

    def strip_wrap(f):
        return f

    # ---- multi-building detail block (buildings 2..12) ----------------------
    r_title, r_note, r_b1depr = 78, 79, 80
    ws[f"B{r_title}"] = "MULTI-BUILDING BASE-COST DETAIL  (Buildings 2-12)"
    ws[f"B{r_title}"].fill = HDR_BLUE
    ws[f"B{r_title}"].font = HDR_FONT
    ws[f"B{r_note}"] = ("Building 1 is the summary above (Steps 1-3 read Property Info "
                        "Building 1). Each column below reads a building from Property "
                        "Info cols D-N and adds its Replacement Cost New to the total.")
    ws[f"B{r_note}"].font = NOTE_FONT

    ws[f"B{r_b1depr}"] = "Building 1 Physical Depr % (auto)"
    ws[f"C{r_b1depr}"] = T["C60"]          # original single-building depr% lookup

    # attribute rows for the buildings-2..12 table
    R = dict(label=82, occ=83, cls=84, qual=85, area=86, ht=87, ap1=88, ap2=89,
             spr=90, age=91, base=92, sprk=93, ccm=94, lam=95, apm=96, shm=97,
             prod=98, adj=99, life=100, dpct=101, ddol=102)
    labels = {
        82: "Building", 83: "Occupancy Type", 84: "Building Class", 85: "Quality",
        86: "Area GBA (SF)", 87: "Avg Eave/Wall Ht (ft)", 88: "Area/Perim Factor 1",
        89: "Area/Perim Factor 2", 90: "Sprinkler Tier", 91: "Effective Age (yrs)",
        92: "Base Cost $/SF (auto)", 93: "Sprinkler $/SF (auto)",
        94: "CCM / De-trend (auto)", 95: "Local Area Mult (auto)",
        96: "Area/Perimeter Mult (auto)", 97: "Story Height Mult (auto)",
        98: "Product of Multipliers (auto)", 99: "Adjusted Base Cost (auto)",
        100: "Typical Life (auto)", 101: "Physical Depr % (auto)",
        102: "Physical Depr $ (auto)",
    }
    for row, lab in labels.items():
        cell = ws[f"B{row}"]
        cell.value = lab
        cell.font = BOLD if row == 82 else Font(size=10)
        if row == 82:
            cell.fill = SUBHDR

    # buildings 2..12 -> cost-approach cols C..M ; Property Info cols D..N
    ca_cols = [get_column_letter(3 + i) for i in range(N_BLDG - 1)]   # C..M
    for i, cc in enumerate(ca_cols):
        bldg = i + 2                                   # building number 2..12
        pcol = get_column_letter(2 + bldg)             # Property Info column D..N

        ws[f"{cc}{R['label']}"] = f"='Property Info'!{pcol}1"
        ws[f"{cc}{R['label']}"].font = BOLD
        ws[f"{cc}{R['label']}"].alignment = Alignment(horizontal="center")
        # raw inputs pulled from Property Info
        ws[f"{cc}{R['occ']}"] = f"='Property Info'!{pcol}35"
        ws[f"{cc}{R['cls']}"] = f"='Property Info'!{pcol}36"
        ws[f"{cc}{R['qual']}"] = f"='Property Info'!{pcol}37"
        ws[f"{cc}{R['area']}"] = f"='Property Info'!{pcol}5"
        ws[f"{cc}{R['ht']}"] = f"='Property Info'!{pcol}38"
        ws[f"{cc}{R['ap1']}"] = f"='Property Info'!{pcol}39"
        ws[f"{cc}{R['ap2']}"] = f"='Property Info'!{pcol}40"
        ws[f"{cc}{R['spr']}"] = f"='Property Info'!{pcol}41"
        ws[f"{cc}{R['age']}"] = f"='Property Info'!{pcol}42"

        occ, cls, qual = f"{cc}{R['occ']}", f"{cc}{R['cls']}", f"{cc}{R['qual']}"
        area, ht, spr, age = (f"{cc}{R['area']}", f"{cc}{R['ht']}",
                              f"{cc}{R['spr']}", f"{cc}{R['age']}")

        # base cost $/SF
        base_core = re.sub(r"^=", "", T["D24"])
        base_core = base_core.rsplit(",", 1)[0]        # drop the ,"NOT FOUND")
        base_core = re.sub(r"^IFERROR\(", "", base_core)
        base_f = sub_refs(base_core, {"$C$4": occ, "$C$6": cls, "$C$7": qual})
        ws[f"{cc}{R['base']}"] = f"=IFERROR({base_f},0)"

        # sprinkler $/SF
        spr_f = sub_refs(re.sub(r"^=", "", T["D25"]),
                         {"$C$4": occ, "$C$24": area, "$C$14": spr})
        ws[f"{cc}{R['sprk']}"] = "=" + spr_f

        # CCM / de-trend  (dates stay shared: $C$17/$C$18/$C$19)
        ccm_f = sub_refs(re.sub(r"^=", "", T["D33"]), {"$C$6": cls, "$C$4": occ})
        ws[f"{cc}{R['ccm']}"] = "=IFERROR(" + ccm_f + ",0)"

        # local area mult (city stays shared: $C$8)
        lam_f = sub_refs(re.sub(r"^=", "", T["D34"]), {"$C$6": cls})
        ws[f"{cc}{R['lam']}"] = "=IFERROR(" + lam_f + ",0)"

        # area/perimeter mult (defaults to 1.0 when factors are left blank)
        ws[f"{cc}{R['apm']}"] = (f"=IF(AVERAGE({cc}{R['ap1']},{cc}{R['ap2']})>0,"
                                 f"AVERAGE({cc}{R['ap1']},{cc}{R['ap2']}),1)")

        # story height mult
        shm_f = sub_refs(re.sub(r"^=", "", T["D36"]), {"$C$4": occ, "$C$11": ht})
        ws[f"{cc}{R['shm']}"] = "=IFERROR(" + shm_f + ",1)"

        # product of multipliers
        ws[f"{cc}{R['prod']}"] = (f"=ROUND(PRODUCT({cc}{R['ccm']},{cc}{R['lam']},"
                                  f"{cc}{R['apm']},{cc}{R['shm']}),3)")

        # adjusted base cost  (only when the building carries an area)
        ws[f"{cc}{R['adj']}"] = (f"=IF({cc}{R['area']}>0,{cc}{R['prod']}*{cc}{R['area']}"
                                 f"*({cc}{R['base']}+{cc}{R['sprk']}),0)")

        # typical life
        life_core = re.sub(r"^=IFERROR\(", "", T["C59"])
        life_core = life_core.rsplit(',"n/a', 1)[0]
        life_f = sub_refs(life_core, {"$C$4": occ, "$C$7": qual, "$C$6": cls})
        ws[f"{cc}{R['life']}"] = f'=IFERROR({life_f},"")'

        # physical depreciation %
        dpct = (f"INDEX(DEPR_VALUES,MATCH({cc}{R['age']},DEPR_AGE,1),"
                f"MATCH({cc}{R['life']},DEPR_LIFEHDR,0))/100")
        ws[f"{cc}{R['dpct']}"] = f"=IFERROR({dpct},0)"

        # physical depreciation $
        ws[f"{cc}{R['ddol']}"] = (f"=IFERROR({cc}{R['adj']}*{cc}{R['dpct']},0)")

    # ---- totals for buildings 2..12 -----------------------------------------
    sum_cols = f"C{{r}}:{ca_cols[-1]}{{r}}"
    ws["B104"] = "Total Adj. Base Cost, Buildings 2-12 (auto)"
    ws["B104"].font = BOLD
    ws["C104"] = f"=SUM({sum_cols.format(r=R['adj'])})"
    ws["B105"] = "Total Physical Depr $, Buildings 2-12 (auto)"
    ws["B105"].font = BOLD
    ws["C105"] = f"=SUM({sum_cols.format(r=R['ddol'])})"

    # ---- wire the totals into the summary -----------------------------------
    # Adjusted Base Building Cost = Building 1 + Buildings 2..12
    ws["G38"] = "=F37*F31+C104"

    # Physical depreciation % blends all buildings, weighted by adjusted base cost
    ws["C60"] = ("=IF('Property Info'!B4>=2,"
                 "IFERROR((F37*F31*C80+C105)/(F37*F31+C104),C80),C80)")
    ws["I60"] = ("MVS Sec.97 percent-good curve. With 2+ buildings this blends each "
                 "building's physical depreciation, weighted by adjusted base cost.")

    # remove the now-obsolete input dropdowns on the mirrored Building-1 cells
    keep = []
    for dv in list(ws.data_validations.dataValidation):
        sq = str(dv.sqref)
        if sq in ("C4", "C6", "C7", "C14"):
            continue
        keep.append(dv)
    ws.data_validations.dataValidation = keep

    # light formatting for the detail block inputs (visual grouping only)
    for cc in ca_cols:
        for row in range(R['occ'], R['age'] + 1):
            ws[f"{cc}{row}"].fill = AUTO_GREY
        for row in range(R['base'], R['ddol'] + 1):
            ws[f"{cc}{row}"].fill = PatternFill(fill_type=None)


# =============================================================================
# 3. BASIC SALES GRID  — total SF + primary use from the main subject
# =============================================================================
def sales_grid(wb):
    ws = wb["Basic Sales Grid"]
    ws["D7"] = "='Property Info'!B16"      # Property Type = primary use
    ws["D13"] = "='Property Info'!B5"      # Property Size (SF) = total GBA
    ws["D29"] = "='Property Info'!B5"      # qualitative size = total GBA
    # D41 already = 'Property Info'!B5


if __name__ == "__main__":
    build()
