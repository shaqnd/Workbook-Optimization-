#!/usr/bin/env python3
"""
1. Import the Adams County zoning directory into the 2025_SALES_GRID workbook
   as a new 'Zoning Directory' sheet.
2. Fix the land-to-building (L:B) ratio on the WO tab so each comp's ratio is
   computed from its own displayed Land / Building Area (matching the subject),
   instead of the inconsistently-calculated Sales!T column.
"""
import warnings
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

warnings.filterwarnings("ignore")

WB = "2025_SALES_GRID.xlsm"
ZONING = ("/root/.claude/uploads/36b651bf-44b6-5554-a780-56e133c5d552/"
          "30136810-adams_county_zoning_directory.xlsx")

HDR_FILL = PatternFill("solid", fgColor="4472C4")
HDR_FONT = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def import_zoning(wb):
    src = openpyxl.load_workbook(ZONING, data_only=True)["Zoning Directory"]
    if "Zoning Directory" in wb.sheetnames:
        del wb["Zoning Directory"]
    ws = wb.create_sheet("Zoning Directory")

    # title banner
    ncols = src.max_column
    ws.cell(1, 1, "ADAMS COUNTY ZONING DIRECTORY  (imported reference)")
    ws.cell(1, 1).font = Font(bold=True, size=12)

    for r in range(1, src.max_row + 1):          # copy the table starting row 3
        for c in range(1, ncols + 1):
            v = src.cell(r, c).value
            dst = ws.cell(r + 2, c, v)
            dst.border = BOX
            if r == 1:                            # header row
                dst.fill = HDR_FILL
                dst.font = HDR_FONT
                dst.alignment = Alignment(horizontal="center", wrap_text=True)
            else:
                dst.alignment = Alignment(vertical="top", wrap_text=True)

    widths = {1: 24, 2: 16, 3: 14, 4: 16, 5: 11, 6: 14, 7: 30, 8: 30,
              9: 34, 10: 14, 11: 40}
    from openpyxl.utils import get_column_letter
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A4"

    # position the sheet next to the existing Zoning Codes tab
    if "Zoning Codes" in wb.sheetnames:
        idx = wb.sheetnames.index("Zoning Codes") + 1
        wb.move_sheet("Zoning Directory",
                      offset=idx - wb.sheetnames.index("Zoning Directory"))
    return src.max_row - 1


def fix_lb_ratio(wb):
    ws = wb["WO"]
    # comp block rows: (land_row, ratio_row, building_row)
    blocks = [(20, 21, 22), (23, 24, 25), (26, 27, 28)]
    for land, ratio, bldg in blocks:
        ws[f"D{ratio}"] = (
            f'=IF(OR(D{bldg}="",D{bldg}=0),"",ROUND(D{land}/D{bldg},1)&":1")'
        )


def main():
    wb = openpyxl.load_workbook(WB, keep_vba=True)
    n = import_zoning(wb)
    fix_lb_ratio(wb)
    try:
        wb.calculation.fullCalcOnLoad = True
    except Exception:
        pass
    wb.save(WB)
    print(f"saved {WB}; imported {n} jurisdictions; fixed WO L:B ratio")


if __name__ == "__main__":
    main()
